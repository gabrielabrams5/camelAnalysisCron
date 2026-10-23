#!/usr/bin/env python3
"""
Parser for the QuickBooks "Statement of Activity by Month - Camel" workbook.

Every Sunday QuickBooks emails "Camel P&L by Month updated weekly" with this
workbook attached. It covers the fiscal year to date (July 1 - last Sunday),
one column per month plus a Total column, laid out like:

    Hillel Foundation of Cambridge
    Statement of Activity by Month - Camel
    July 1-October 18, 2026
                                   Jul 2026  ...  Oct 1-18 2026  Total
    Income
      425040 The Camel                ...                         ...
    Total Income                                                  ...
    Expenses
      610030 Food and Beverage        ...                         ...
    Total Expenses                                                ...
    Net Income                                                    ...

Older copies say "Revenue"/"Expenditures"/"Net Revenue" instead of
"Income"/"Expenses"/"Net Income"; both are accepted.

QuickBooks writes a total cell as a plain number when it is nonzero and as a
formula over blank cells (e.g. "=C8+C9") when it is zero, without a cached
result. So formulas are evaluated here (sums of cell references only) instead
of trusting openpyxl's data_only values, which come back as None for those.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Dict, List, Optional

import openpyxl

CENT = Decimal('0.01')

TITLE_MARKERS = ('statement of activity', 'profit and loss')

_REF_OR_NUM = r'(?:\$?[A-Z]{1,3}\$?\d+|\d+(?:\.\d+)?)'
_FORMULA = re.compile(rf'^[+-]?{_REF_OR_NUM}(?:[+-]{_REF_OR_NUM})*$')
_TERM = re.compile(rf'([+-]?)({_REF_OR_NUM})')

# "July 1-October 18, 2026", "July 1-5, 2027", "July 1, 2026-January 3, 2027"
_PERIOD = re.compile(
    r'([A-Z][a-z]+) (\d{1,2})(?:, (\d{4}))?\s*[-–]\s*'
    r'(?:([A-Z][a-z]+) )?(\d{1,2}), (\d{4})'
)

SECTION_HEADERS = {
    'income': 'income', 'revenue': 'income',
    'expenses': 'expense', 'expenditures': 'expense',
    'other income': 'other_income', 'other revenue': 'other_income',
    'other expenses': 'other_expense', 'other expenditures': 'other_expense',
}

TOTAL_ROWS = {
    'income': re.compile(r'^total (for )?(income|revenue)$'),
    'expense': re.compile(r'^total (for )?(expenses|expenditures)$'),
    'other_income': re.compile(r'^total (for )?other (income|revenue)$'),
    'other_expense': re.compile(r'^total (for )?other (expenses|expenditures)$'),
    'net': re.compile(r'^net (income|revenue)$'),
}

# Leaf account rows start with the QuickBooks account number: "610030 Food and Beverage"
_ACCOUNT = re.compile(r'^\d{4,}\s')


class ReportFormatError(ValueError):
    """The workbook looks like a P&L but could not be read."""


@dataclass
class PLReport:
    period_start: date
    as_of: date
    total_income: Decimal
    total_expenses: Decimal
    net_income: Decimal
    # [{"account", "section", "months": {"Jul 2026": 123.45, ...}, "total"}]
    accounts: List[Dict] = field(default_factory=list)

    @property
    def fiscal_year(self) -> int:
        return fiscal_year_of(self.as_of)


def fiscal_year_of(d: date) -> int:
    """Hillel's fiscal year runs July-June and is named for the year it ends in
    (July 2026 - June 2027 is FY27)."""
    return d.year + 1 if d.month >= 7 else d.year


def fy_label(fiscal_year: int) -> str:
    return f"FY{fiscal_year % 100:02d}"


def parse_pl_workbook(data: bytes) -> Optional[PLReport]:
    """
    Parse an emailed QuickBooks workbook.

    Returns None if no sheet in it is a Statement of Activity / P&L (e.g. the
    weekly transaction report QuickBooks also sends). Raises ReportFormatError
    if there is one but it cannot be read.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = _find_pl_sheet(wb)
    if ws is None:
        return None
    return _parse_sheet(ws)


def _find_pl_sheet(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=6, max_col=3, values_only=True):
            for value in row:
                if isinstance(value, str) and any(m in value.lower() for m in TITLE_MARKERS):
                    return ws
    return None


def _parse_sheet(ws) -> PLReport:
    period_start, as_of = _parse_period(ws)
    header_row, amount_cols, total_col = _find_header(ws)

    totals: Dict[str, Decimal] = {}
    accounts = []
    section = None

    for r in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not isinstance(label, str) or not label.strip():
            continue
        label = label.strip()
        key = re.sub(r'\s+', ' ', label.lower())

        if key in SECTION_HEADERS:
            section = SECTION_HEADERS[key]
            continue

        matched_total = next((name for name, rx in TOTAL_ROWS.items() if rx.match(key)), None)
        if matched_total:
            totals[matched_total] = _amount(ws, ws.cell(row=r, column=total_col).value)
            continue

        if _ACCOUNT.match(label) and not key.startswith('total'):
            months = {
                header: _amount(ws, ws.cell(row=r, column=c).value)
                for c, header in amount_cols
            }
            # Parent accounts ("6100 Consumable Supplies & Expenses") are bare
            # headings with no amounts; only rows that carry money are kept.
            if all(ws.cell(row=r, column=c).value in (None, '') for c, _ in amount_cols) \
                    and ws.cell(row=r, column=total_col).value in (None, ''):
                continue
            accounts.append({
                'account': label,
                'section': section,
                'months': {h: float(v) for h, v in months.items()},
                'total': float(_amount(ws, ws.cell(row=r, column=total_col).value)),
            })

    if not totals:
        raise ReportFormatError(
            f"sheet '{ws.title}': found no Total Income / Total Expenses / Net Income rows")

    income = totals.get('income', Decimal('0')) + totals.get('other_income', Decimal('0'))
    expenses = totals.get('expense', Decimal('0')) + totals.get('other_expense', Decimal('0'))
    net = totals.get('net', income - expenses)
    if abs(income - expenses - net) > CENT:
        raise ReportFormatError(
            f"sheet '{ws.title}': income {income} - expenses {expenses} != net {net}; "
            "the layout may have changed")

    return PLReport(
        period_start=period_start,
        as_of=as_of,
        total_income=income.quantize(CENT),
        total_expenses=expenses.quantize(CENT),
        net_income=net.quantize(CENT),
        accounts=accounts,
    )


def _parse_period(ws):
    for row in ws.iter_rows(min_row=1, max_row=6, max_col=3, values_only=True):
        for value in row:
            if not isinstance(value, str):
                continue
            m = _PERIOD.search(value)
            if not m:
                continue
            start_month, start_day, start_year, end_month, end_day, end_year = m.groups()
            end_month = end_month or start_month
            try:
                end = datetime.strptime(f"{end_month} {end_day} {end_year}", '%B %d %Y').date()
                start = datetime.strptime(
                    f"{start_month} {start_day} {start_year or end.year}", '%B %d %Y').date()
            except ValueError:
                continue
            if start > end:  # "July 1-January 3, 2027": the start is in the prior year
                start = start.replace(year=start.year - 1)
            return start, end
    raise ReportFormatError(f"sheet '{ws.title}': no reporting period like 'July 1-October 18, 2026'")


def _find_header(ws):
    """The header row is the one with a 'Total' column; month columns sit left of it."""
    for r in range(1, min(ws.max_row, 12) + 1):
        cells = [(c, ws.cell(row=r, column=c).value) for c in range(2, ws.max_column + 1)]
        labels = [(c, v.strip()) for c, v in cells if isinstance(v, str) and v.strip()]
        if not labels:
            continue
        total_cols = [c for c, v in labels if v.lower() == 'total']
        if total_cols:
            total_col = total_cols[0]
            return r, [(c, v) for c, v in labels if c < total_col], total_col
        # A single-month report may come without a Total column
        if r > 3 and len(labels) == 1:
            c, v = labels[0]
            return r, [(c, v)], c
    raise ReportFormatError(f"sheet '{ws.title}': no header row with a 'Total' column")


def _amount(ws, value, depth=0) -> Decimal:
    if value is None or value == '':
        return Decimal('0')
    if isinstance(value, bool):
        raise ReportFormatError(f"unexpected boolean in an amount cell")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, str) and value.startswith('='):
        expr = value[1:].replace(' ', '')
        if not _FORMULA.match(expr):
            raise ReportFormatError(f"unsupported formula {value!r}")
        if depth > 25:
            raise ReportFormatError(f"formula {value!r} nests too deep")
        total = Decimal('0')
        for sign, term in _TERM.findall(expr):
            if term[0].isdigit():
                part = Decimal(term)
            else:
                part = _amount(ws, ws[term.replace('$', '')].value, depth + 1)
            total += -part if sign == '-' else part
        return total
    raise ReportFormatError(f"unexpected text {value!r} in an amount cell")


if __name__ == '__main__':
    import sys
    for path in sys.argv[1:]:
        with open(path, 'rb') as fh:
            report = parse_pl_workbook(fh.read())
        if report is None:
            print(f"{path}: no Statement of Activity sheet")
            continue
        print(f"{path}: {fy_label(report.fiscal_year)} {report.period_start} -> {report.as_of}  "
              f"income {report.total_income}  expenses {report.total_expenses}  "
              f"net {report.net_income}  ({len(report.accounts)} accounts)")
